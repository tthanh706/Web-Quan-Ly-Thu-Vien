import sys
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

def create_speech_word_doc():
    doc = Document()

    # Cấu hình lề trang 2.54 cm (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    FONT_FAMILY = 'Times New Roman'
    COLOR_BLACK = RGBColor(0, 0, 0)

    # --- TIÊU ĐỀ TRANG ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(18)
    
    r_sub = title_p.add_run("KỊCH BẢN THUYẾT TRÌNH BÁO CÁO TIẾN ĐỘ DỰ ÁN (PHIÊN BẢN NGẮN GỌN 3 - 5 PHÚT)\n")
    r_sub.font.name = FONT_FAMILY
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_BLACK

    r_main = title_p.add_run("LỜI NÓI THUYẾT TRÌNH TỪNG SLIDE (10 SLIDES)\nHỆ THỐNG QUẢN LÝ THƯ VIỆN")
    r_main.font.name = FONT_FAMILY
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = COLOR_BLACK

    # Thông tin chung
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.space_after = Pt(24)
    
    r_info = info_p.add_run(
        "Môn học: Phân Tích & Thiết Kế Hệ Thống Thông Tin / Lập Trình Web\n"
        "Đề tài: Hệ thống Quản lý Thư viện (SmartLibrary)\n"
        "Nhóm 3 - Lớp Quản Lý Thư Viện (Thành viên: Trương Thị Hạnh, Nguyễn Hồng Ngọc, Vy Thị Hồng Ngọc)"
    )
    r_info.font.name = FONT_FAMILY
    r_info.font.size = Pt(11.5)
    r_info.font.italic = True
    r_info.font.color.rgb = COLOR_BLACK

    def add_slide_header(slide_num, title, duration):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        
        run_num = h.add_run(f"SLIDE {slide_num}: {title.upper()}\n")
        run_num.font.name = FONT_FAMILY
        run_num.font.size = Pt(12.5)
        run_num.font.bold = True
        run_num.font.color.rgb = COLOR_BLACK

        run_dur = h.add_run(f"⏱️ Thời lượng phát biểu dự kiến: {duration}")
        run_dur.font.name = FONT_FAMILY
        run_dur.font.size = Pt(10.5)
        run_dur.font.italic = True
        run_dur.font.color.rgb = COLOR_BLACK
        return h

    def add_speech_box(speech_text, notes_text=""):
        p_speech = doc.add_paragraph()
        p_speech.paragraph_format.left_indent = Inches(0.3)
        p_speech.paragraph_format.right_indent = Inches(0.3)
        p_speech.paragraph_format.space_after = Pt(4)
        p_speech.paragraph_format.line_spacing = 1.2

        r_lbl = p_speech.add_run("🗣️ Lời nói của Sinh viên (Đọc/Nói trực tiếp):\n")
        r_lbl.font.name = FONT_FAMILY
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_BLACK

        r_txt = p_speech.add_run(f'"{speech_text}"')
        r_txt.font.name = FONT_FAMILY
        r_txt.font.size = Pt(12)
        r_txt.font.italic = True
        r_txt.font.color.rgb = COLOR_BLACK

        if notes_text:
            p_notes = doc.add_paragraph()
            p_notes.paragraph_format.left_indent = Inches(0.3)
            p_notes.paragraph_format.space_after = Pt(10)
            
            r_nlbl = p_notes.add_run("💡 Mẹo trình diễn & Thao tác trên slide: ")
            r_nlbl.font.name = FONT_FAMILY
            r_nlbl.font.size = Pt(10.5)
            r_nlbl.font.bold = True
            r_nlbl.font.color.rgb = COLOR_BLACK

            r_ntxt = p_notes.add_run(notes_text)
            r_ntxt.font.name = FONT_FAMILY
            r_ntxt.font.size = Pt(10.5)
            r_ntxt.font.color.rgb = COLOR_BLACK

    # --- NỘI DUNG LỜI THUYẾT TRÌNH TỪNG SLIDE (BẢN NGẮN GỌN 3-5 PHÚT) ---

    # Slide 1
    add_slide_header(1, "Thông tin Đề tài & Nhóm thực hiện", "Khoảng 15 giây")
    add_speech_box(
        "Kính chào Thầy/Cô và các bạn. Em xin đại diện Nhóm 3 trình bày báo cáo tiến độ dự án: 'Hệ thống Quản lý Thư viện SmartLibrary'. Nhóm em gồm 3 thành viên: Trương Thị Hạnh (Backend & CSDL), Nguyễn Hồng Ngọc (Frontend SPA) và Vy Thị Hồng Ngọc (Phân tích & Kiểm thử).",
        "Chào Thầy/Cô với thái độ tự tin, mỉm cười nhẹ và chỉ tay vào tên từng thành viên trên Slide 1."
    )

    # Slide 2
    add_slide_header(2, "Giới thiệu Đề tài & Vấn đề cần giải quyết", "Khoảng 20 giây")
    add_speech_box(
        "Việc quản lý thư viện truyền thống thủ công hiện gặp 2 bất cập lớn: độc giả mất nhiều thời gian tìm vị trí sách, còn thủ thư dễ ghi nhầm hạn mượn và tính sai tiền phạt. Vì vậy, nhóm em xây dựng Web SmartLibrary nhằm tự động hóa hoàn toàn các nghiệp vụ này.",
        "Nhấn giọng ngắn gọn ở 2 cụm từ: 'mất thời gian' và 'tính sai tiền phạt'."
    )

    # Slide 3
    add_slide_header(3, "Mục tiêu Hệ thống, Đối tượng & Phạm vi", "Khoảng 20 giây")
    add_speech_box(
        "Mục tiêu cốt lõi của project là tạo ra hệ thống quản lý trực tuyến 24/7, hỗ trợ mượn/trả tự động và tự tính phạt quá hạn 5.000đ/ngày. Hệ thống phục vụ 3 nhóm người dùng: Quản trị viên, Thủ thư và Độc giả.",
        "Chỉ nhanh vào 3 icon biểu tượng 3 vai trò người dùng trên Slide 3."
    )

    # Slide 4
    add_slide_header(4, "Phân tích Yêu cầu Chức năng & Phi chức năng", "Khoảng 25 giây")
    add_speech_box(
        "Về yêu cầu: Hệ thống đáp ứng phân quyền 3 vai trò, tra cứu sách real-time, quản lý thẻ độc giả, lập phiếu mượn 14 ngày, gia hạn, trả sách và đặt trước. Giao diện được thiết kế tươi sáng, phản hồi mượt màng dưới 1 giây và bảo mật mật khẩu bằng SHA-256.",
        "Nói nhịp điệu dứt khoát, mượt mà."
    )

    # Slide 5
    add_slide_header(5, "Các Chức năng chính thực sự có trong Project", "Khoảng 25 giây")
    add_speech_box(
        "Trực quan trên Slide 5 là 6 cụm chức năng chính đã được nhóm em lập trình hoàn thiện 100%: Đăng nhập phân quyền, Tra cứu kho sách & kệ kho, Quản lý thẻ độc giả, Quy trình mượn/trả/gia hạn/thu phạt, Đặt trước sách và Biểu đồ báo cáo Dashboard.",
        "Khẳng định rõ ràng: Toàn bộ 6 chức năng đều đã hoàn thành thực tế và sẵn sàng demo."
    )

    # Slide 6
    add_slide_header(6, "Sơ đồ Chức năng & Các Actor trong Hệ thống", "Khoảng 25 giây")
    add_speech_box(
        "Về phân quyền actor: Admin quản lý người dùng và xem báo cáo tổng quan. Thủ thư trực tiếp quản lý kho sách, cấp thẻ và xử lý phiếu mượn/trả tại quầy. Độc giả có thể tra cứu vị trí kệ sách và theo dõi hạn trả phiếu mượn cá nhân.",
        "Chỉ tay lần lượt qua 3 cột Admin -> Thủ thư -> Độc giả trên Slide 6."
    )

    # Slide 7
    add_slide_header(7, "Thiết kế Cơ sở Dữ liệu 7 Bảng (Visual ERD)", "Khoảng 25 giây")
    add_speech_box(
        "Cơ sở dữ liệu của hệ thống được chuẩn hóa gồm 7 bảng trên SQLite3: Bảng users, readers, categories, books, borrow_records, reservations và fine_logs. Các bảng liên kết chặt chẽ qua khóa ngoại reader_id, category_id và book_id.",
        "Nhìn nhanh vào hình vẽ sơ đồ ERD 7 bảng trên Slide 7."
    )

    # Slide 8
    add_slide_header(8, "Kiến trúc Hệ thống & Mô hình Giao tiếp REST API", "Khoảng 25 giây")
    add_speech_box(
        "Hệ thống theo mô hình RESTful Client-Server decouple: Frontend là Web SPA HTML/CSS/JS giao tiếp với Backend Python 8000 & SQLite3 hoàn toàn qua dữ liệu chuẩn JSON, giúp trang web hoạt động mượt màng không cần tải lại trang.",
        "Nhấn mạnh tính mượt mà của Web Single Page Application (SPA)."
    )

    # Slide 9
    add_slide_header(9, "Web Demo & Quy trình trình diễn đề xuất", "Khoảng 30 giây")
    add_speech_box(
        "Để chứng minh thực tế, nhóm em chuẩn bị 7 bước demo live: Đăng nhập Thủ thư -> Tra cứu sách -> Thêm sách mới vào kho -> Tạo thẻ độc giả -> Lập phiếu mượn -> Trả sách & Thu tiền phạt quá hạn -> Và xem Báo cáo Dashboard tự cập nhật.",
        "Chuyển giọng hào hứng, chuẩn bị thao tác demo trực tiếp trên trình duyệt Web."
    )

    # Slide 10
    add_slide_header(10, "Kết quả đạt được, Hạn chế & Hướng phát triển", "Khoảng 20 giây")
    add_speech_box(
        "Tổng kết lại, nhóm đã hoàn thành 100% ứng dụng Web mượt mà. Hướng phát triển tiếp theo là tích hợp Camera quét mã Barcode và Mobile App thông báo. Em xin chân thành cảm ơn Thầy/Cô và các bạn đã lắng nghe!",
        "Cúi đầu chào cảm ơn lịch sự khi hoàn tất bài thuyết trình."
    )

    output_path = "Kich_Ban_Thuyet_Trinh_10_Slide.docx"
    try:
        doc.save(output_path)
    except Exception:
        output_path = "Kich_Ban_Thuyet_Trinh_10_Slide_v2.docx"
        doc.save(output_path)
    print(f"Successfully generated Speech Word document at: {os.path.abspath(output_path)}")

if __name__ == '__main__':
    create_speech_word_doc()
