import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

def create_10_slide_presentation():
    prs = Presentation()
    # Kích thước slide chuẩn 16:9 Widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- BẢNG MÀU TƯƠI SÁNG & HIỆN ĐẠI (Bright Light Palette) ---
    COLOR_BG = RGBColor(248, 250, 252)          # Nền sáng mềm mại (#F8FAFC)
    COLOR_CARD = RGBColor(255, 255, 255)        # Thẻ trắng tinh tế (#FFFFFF)
    COLOR_CARD_BORDER = RGBColor(226, 232, 240)  # Viền thẻ xám nhẹ (#E2E8F0)
    
    # Màu sắc chủ đạo & điểm nhấn
    COLOR_PRIMARY = RGBColor(79, 70, 229)      # Xanh Indigo hiện đại (#4F46E5)
    COLOR_CYAN = RGBColor(14, 165, 233)         # Xanh da trời tươi sáng (#0EA5E9)
    COLOR_EMERALD = RGBColor(16, 185, 129)     # Xanh lá ngọc tươi (#10B981)
    COLOR_ROSE = RGBColor(244, 63, 94)         # Hồng đỏ nổi bật (#F43F5E)
    COLOR_AMBER = RGBColor(245, 158, 11)       # Vàng cam ấm áp (#F59E0B)

    # Chữ tối màu sắc nét trên nền sáng
    COLOR_TEXT_MAIN = RGBColor(15, 23, 42)      # Chữ chính màu xanh đen đậm (#0F172A)
    COLOR_TEXT_MUTED = RGBColor(71, 85, 105)    # Chữ phụ màu xám đậm (#475569)

    def set_background(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, category_text="BÁO CÁO TIẾN ĐỘ DỰ ÁN - HỆ THỐNG QUẢN LÝ THƯ VIỆN"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = 'Segoe UI'
        p0.font.size = Pt(10.5)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_CYAN

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = 'Segoe UI'
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY

    def add_card(slide, left, top, width, height, title, content_bullets, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, title_color=COLOR_PRIMARY, font_size=Pt(11.5)):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = 'Segoe UI'
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = title_color
        p0.space_after = Pt(8)

        for b in content_bullets:
            p = tf.add_paragraph()
            if isinstance(b, tuple):
                p.text = b[0]
                p.font.bold = b[1]
                p.font.size = font_size
                p.font.color.rgb = b[2] if len(b) > 2 else COLOR_TEXT_MAIN
            else:
                p.text = "• " + b
                p.font.size = font_size
                p.font.color.rgb = COLOR_TEXT_MAIN
            p.font.name = 'Segoe UI'
            p.space_after = Pt(4.5)

    def draw_erd_entity(slide, left, top, width, height, title_text, fields, header_bg=COLOR_PRIMARY, border_color=COLOR_CARD_BORDER):
        """Vẽ một khối Bảng CSDL (ERD Entity Table) chuyên nghiệp có Banner màu và danh sách Trường dữ liệu"""
        # Thẻ container chính
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        # Banner Tiêu đề Bảng
        header_h = 0.42
        hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(header_h))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = header_bg
        hdr.line.fill.background()

        tf_hdr = hdr.text_frame
        tf_hdr.word_wrap = True
        tf_hdr.margin_left = Inches(0.12)
        tf_hdr.margin_top = Inches(0.06)
        p_hdr = tf_hdr.paragraphs[0]
        p_hdr.text = title_text
        p_hdr.font.name = 'Segoe UI'
        p_hdr.font.size = Pt(11.5)
        p_hdr.font.bold = True
        p_hdr.font.color.rgb = RGBColor(255, 255, 255)

        # Danh sách trường dữ liệu
        tb_body = slide.shapes.add_textbox(Inches(left + 0.1), Inches(top + header_h + 0.05), Inches(width - 0.2), Inches(height - header_h - 0.1))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        tf_body.margin_left = tf_body.margin_top = tf_body.margin_right = tf_body.margin_bottom = 0

        for idx, (f_name, f_type, is_pk, is_fk) in enumerate(fields):
            p = tf_body.paragraphs[0] if idx == 0 else tf_body.add_paragraph()
            p.space_after = Pt(2.5)

            # Icon PK/FK
            if is_pk and is_fk:
                prefix = "🔑🔗 "
            elif is_pk:
                prefix = "🔑 "
            elif is_fk:
                prefix = "🔗 "
            else:
                prefix = "▫️ "

            r_pre = p.add_run()
            r_pre.text = prefix
            r_pre.font.name = 'Segoe UI'
            r_pre.font.size = Pt(9.5)

            r_name = p.add_run()
            r_name.text = f_name
            r_name.font.name = 'Segoe UI'
            r_name.font.size = Pt(10)
            r_name.font.bold = is_pk or is_fk
            r_name.font.color.rgb = COLOR_PRIMARY if (is_pk or is_fk) else COLOR_TEXT_MAIN

            r_type = p.add_run()
            r_type.text = f"  ({f_type})"
            r_type.font.name = 'Segoe UI'
            r_type.font.size = Pt(9)
            r_type.font.color.rgb = COLOR_TEXT_MUTED

    # ----------------------------------------------------
    # SLIDE 1: Thông tin Đề tài & Nhóm thực hiện
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_background(s1)

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_CARD
    card1.line.color.rgb = COLOR_PRIMARY
    card1.line.width = Pt(2.5)

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(0.9), Inches(10.933), Inches(5.7))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "BÁO CÁO TIẾN ĐỘ DỰ ÁN XÂY DỰNG WEB"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p = tf1.add_paragraph()
    p.text = "HỆ THỐNG QUẢN LÝ THƯ VIỆN"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Môn học: Phân tích & Thiết kế Hệ thống Thông tin / Lập trình Web"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_after = Pt(20)

    p = tf1.add_paragraph()
    p.text = "THÔNG TIN NHÓM & THÀNH VIÊN THỰC HIỆN:"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER
    p.space_after = Pt(8)

    members_detail = [
        "Tên Nhóm: Nhóm 3 - Lớp Quản Lý Thư Viện",
        "Thành viên 1: Trương Thị Hạnh - Lập trình Backend (Python HTTP Server) & CSDL (SQLite3)",
        "Thành viên 2: Nguyễn Hồng Ngọc - Thiết kế UI/UX Giao diện Web (HTML/CSS/JS) & Visual Dashboard",
        "Thành viên 3: Vy Thị Hồng Ngọc - Phân tích Hệ thống (RBAC, Use Case, ERD), QA & Tài liệu"
    ]
    for m in members_detail:
        p = tf1.add_paragraph()
        p.text = "• " + m
        p.font.name = 'Segoe UI'
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(6)

    # ----------------------------------------------------
    # SLIDE 2: Giới thiệu & Lý do chọn đề tài
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_background(s2)
    add_header(s2, "Slide 2: Giới Thiệu Đề Tài & Vấn Đề Cần Giải Quyết")

    add_card(s2, 0.8, 1.5, 5.7, 5.5, "📌 Giới Thiệu & Lý Do Chọn Đề Tài", [
        "Thư viện đóng vai trò trung tâm trong quản lý và lưu trữ tri thức nhà trường và doanh nghiệp.",
        "Nhiều thư viện hiện nay vẫn sử dụng quy trình ghi chép thủ công hoặc phần mềm cũ khó tra cứu.",
        "Xu hướng chuyển đổi số yêu cầu một hệ thống Web quản lý tập trung, minh bạch và hiện đại.",
        "Đề tài hướng tới xây dựng giải pháp Web Thư viện mượt mà, giúp tối ưu hóa công tác mượn/trả sách."
    ], title_color=COLOR_PRIMARY, font_size=Pt(12))

    add_card(s2, 6.8, 1.5, 5.7, 5.5, "❌ Vấn Đề Cần Giải Quyết Thực Tế", [
        "Tra cứu sách lâu: Độc giả mất nhiều thời gian tìm sách, không biết chính xác vị trí kệ kho.",
        "Nhầm lẫn ngày hạn: Thủ thư ghi chép tay dễ nhầm hạn trả và tính sai tiền phạt quá hạn.",
        "Quản lý kho kém: Khó nắm bắt chính xác số lượng sách còn khả dụng thực tế trong kho.",
        "Không có hàng chờ: Khi sách hot bị mượn hết, không có cách nào xếp hàng đặt trước cho người mượn sau."
    ], title_color=COLOR_ROSE, border_color=RGBColor(254, 205, 211), font_size=Pt(12))

    # ----------------------------------------------------
    # SLIDE 3: Mục tiêu & Đối tượng & Phạm vi
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_background(s3)
    add_header(s3, "Slide 3: Mục Tiêu Hệ Thống, Đối Tượng & Phạm Vi")

    add_card(s3, 0.8, 1.5, 3.7, 5.5, "🎯 Mục Tiêu Hệ Thống", [
        "Xây dựng ứng dụng Web SPA tra cứu và quản lý thư viện trực tuyến 24/7.",
        "Số hóa toàn bộ dữ liệu: Sách, Thể loại, Thẻ độc giả, Phiếu mượn/trả.",
        "Tự động hóa hoàn toàn việc tính tiền phạt quá hạn 5.000đ/ngày.",
        "Tối ưu quy trình vận hành cho thủ thư và trải nghiệm tra cứu cho độc giả."
    ], title_color=COLOR_CYAN, border_color=RGBColor(186, 230, 253), font_size=Pt(11.5))

    add_card(s3, 4.8, 1.5, 3.7, 5.5, "👥 Đối Tượng Sử Dụng", [
        "Quản trị viên (Admin): Quản lý tài khoản, phân quyền, xem thống kê báo cáo.",
        "Thủ thư (Librarian): Quản lý sách, quản lý độc giả, cho mượn/trả, gia hạn, thu phạt.",
        "Độc giả (Reader): Tra cứu sách, tìm vị trí kệ, xem lịch sử mượn, đặt trước sách."
    ], title_color=COLOR_PRIMARY, border_color=RGBColor(199, 210, 254), font_size=Pt(11.5))

    add_card(s3, 8.8, 1.5, 3.7, 5.5, "🚩 Phạm Vi Dự Án", [
        "Phạm vi chức năng: Quản lý sách, độc giả, mượn/trả, gia hạn, phạt quá hạn, đặt trước, thống kê.",
        "Phạm vi công nghệ: Ứng dụng Web chạy trên môi trường local/network với Python & SQLite3.",
        "Đối tượng áp dụng: Thư viện trường học, trung tâm đào tạo hoặc thư viện nội bộ."
    ], title_color=COLOR_AMBER, border_color=RGBColor(254, 235, 200), font_size=Pt(11.5))

    # ----------------------------------------------------
    # SLIDE 4: Phân tích Yêu cầu Chức năng & Phi chức năng
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_background(s4)
    add_header(s4, "Slide 4: Phân Tích Yêu Cầu Chức Năng & Phi Chức Năng")

    add_card(s4, 0.8, 1.5, 5.7, 5.5, "⚙️ Yêu Cầu Chức Năng (Functional)", [
        "Đăng nhập & Phân quyền: Đăng nhập bảo mật, phân quyền 3 vai trò (Admin, Thủ thư, Độc giả).",
        "Quản lý Sách & Thể loại: Thêm/sửa/xóa sách, xem vị trí kệ, theo dõi tổng số lượng & tồn kho.",
        "Quản lý Độc giả: Cấp mới thẻ độc giả (Sinh viên, Giảng viên), gia hạn thẻ, tạm khóa thẻ.",
        "Quản lý Mượn/Trả/Phạt: Lập phiếu mượn 14 ngày, gia hạn +14 ngày (tối đa 2 lần), tính phạt 5.000đ/ngày trễ & thu phạt.",
        "Đặt trước sách: Xếp hàng chờ mượn khi sách trong kho hết.",
        "Báo cáo thống kê: Thống kê cơ cấu thể loại, lượt mượn/trả & KPIs realtime."
    ], title_color=COLOR_PRIMARY, border_color=RGBColor(199, 210, 254), font_size=Pt(11))

    add_card(s4, 6.8, 1.5, 5.7, 5.5, "🛡️ Yêu Cầu Phi Chức Năng (Non-Functional)", [
        "Hiệu năng & Phản hồi: Giao diện SPA phản hồi nhanh dưới 1 giây, tra cứu tức thì thời gian thực.",
        "Bảo mật: Mã hóa mật khẩu người dùng bằng SHA-256, kiểm tra session đăng nhập.",
        "Toàn vẹn dữ liệu: CSDL SQLite3 bật ràng buộc khóa ngoại (PRAGMA foreign_keys = ON).",
        "Tính dễ sử dụng: Giao diện Light Mode tươi sáng, trực quan, hỗ trợ nút đăng nhập nhanh dùng thử.",
        "Khả năng mở rộng: Kiến trúc RESTful API cho phép dễ dàng tích hợp thêm tính năng nâng cao."
    ], title_color=COLOR_EMERALD, border_color=RGBColor(167, 243, 208), font_size=Pt(11))

    # ----------------------------------------------------
    # SLIDE 5: Các Chức Năng Chính Thực Sự Có Trong Project
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_background(s5)
    add_header(s5, "Slide 5: Các Chức Năng Chính Thực Sự Có Trong Project")

    add_card(s5, 0.8, 1.5, 5.7, 2.6, "🔐 1. Đăng Nhập & Phân Quyền", [
        "Cổng đăng nhập bắt buộc (Login Gate Overlay).",
        "Phân quyền 3 vai trò: Admin, Thủ thư, Độc giả.",
        "Nút đăng nhập nhanh Preset Accounts dùng thử."
    ], title_color=COLOR_PRIMARY, font_size=Pt(11))

    add_card(s5, 6.8, 1.5, 5.7, 2.6, "📚 2. Tra Cứu & Quản Lý Sách", [
        "Xem sách dạng Thẻ Grid / Bảng Table.",
        "Tìm kiếm từ khóa tức thì & Lọc theo Thể loại.",
        "Thêm/sửa/xóa sách, vị trí kệ (Kệ A1, B2...) & tồn kho."
    ], title_color=COLOR_CYAN, font_size=Pt(11))

    add_card(s5, 0.8, 4.3, 5.7, 2.7, "🤝 3. Quản Lý Mượn / Trả / Phạt", [
        "Lập phiếu mượn (hạn 14 ngày, kiểm tra tồn kho).",
        "Gia hạn hạn mượn +14 ngày (tối đa 2 lần).",
        "Trả sách, tính phạt quá hạn (5.000đ/ngày) & Thu phạt."
    ], title_color=COLOR_ROSE, font_size=Pt(11))

    add_card(s5, 6.8, 4.3, 5.7, 2.7, "🎴 4. Độc Giả, Đặt Trước & Thống Kê", [
        "Quản lý thẻ độc giả (Sinh viên, Giảng viên), khóa/mở thẻ.",
        "Đặt trước sách khi hết kho (Xếp hàng queue_order).",
        "Dashboard Biểu đồ Chart.js & Bảng điều khiển Admin."
    ], title_color=COLOR_EMERALD, font_size=Pt(11))

    # ----------------------------------------------------
    # SLIDE 6: Sơ đồ chức năng / Use Case (Actors)
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_background(s6)
    add_header(s6, "Slide 6: Sơ Đồ Chức Năng & Các Actor Trong Hệ Thống")

    add_card(s6, 0.8, 1.5, 3.7, 5.5, "👑 Actor: Admin (Quản trị)", [
        "• Quản lý tài khoản hệ thống (Users)",
        "• Phân quyền tài khoản (Admin, Thủ thư, Độc giả)",
        "• Xem Dashboard báo cáo tổng quan",
        "• Reset mật khẩu & Cấu hình hệ thống",
        "• Xem Nhật ký nộp phạt (Fine Logs)"
    ], title_color=COLOR_AMBER, border_color=RGBColor(254, 235, 200), font_size=Pt(11.5))

    add_card(s6, 4.8, 1.5, 3.7, 5.5, "📚 Actor: Thủ Thư (Librarian)", [
        "• Quản lý kho sách (Thêm/Sửa/Xóa/Vị trí kệ)",
        "• Quản lý thẻ độc giả (Cấp thẻ, Khóa/Mở thẻ)",
        "• Lập phiếu mượn sách mới tại quầy",
        "• Xác nhận trả sách & Giải phóng tồn kho",
        "• Gia hạn sách & Thu tiền phạt quá hạn",
        "• Duyệt danh sách Đặt trước sách"
    ], title_color=COLOR_PRIMARY, border_color=RGBColor(199, 210, 254), font_size=Pt(11.5))

    add_card(s6, 8.8, 1.5, 3.7, 5.5, "📖 Actor: Độc Giả (Reader)", [
        "• Đăng nhập tài khoản độc giả",
        "• Tra cứu danh mục sách & vị trí kệ kho",
        "• Xem chi tiết thông tin sách & số lượng còn",
        "• Xem lịch sử phiếu mượn cá nhân & hạn trả",
        "• Thực hiện Đặt trước khi sách bị mượn hết"
    ], title_color=COLOR_EMERALD, border_color=RGBColor(167, 243, 208), font_size=Pt(11.5))

    # ----------------------------------------------------
    # SLIDE 7: SƠ ĐỒ CƠ SỞ DỮ LIỆU GRAPHICAL ERD DIAGRAM (VẼ TRỰC QUAN)
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_background(s7)
    add_header(s7, "Slide 7: Sơ Đồ Cơ Sở Dữ Liệu 7 Bảng (Visual ERD Diagram)")

    # Dòng trên (3 Bảng CSDL)
    draw_erd_entity(s7, 0.8, 1.5, 3.7, 2.3, "👤 USERS (Tài khoản)", [
        ("id", "INTEGER", True, False),
        ("username", "TEXT", False, False),
        ("password_hash", "TEXT", False, False),
        ("role", "TEXT", False, False),
        ("full_name", "TEXT", False, False),
        ("reader_id", "INTEGER", False, True)
    ], header_bg=COLOR_PRIMARY)

    draw_erd_entity(s7, 4.8, 1.5, 3.7, 2.3, "🎴 READERS (Độc giả)", [
        ("id", "INTEGER", True, False),
        ("reader_code", "TEXT", False, False),
        ("full_name", "TEXT", False, False),
        ("email / phone", "TEXT", False, False),
        ("card_type", "TEXT", False, False),
        ("expiry_date", "DATE", False, False)
    ], header_bg=COLOR_EMERALD)

    draw_erd_entity(s7, 8.8, 1.5, 3.7, 2.3, "🏷️ CATEGORIES (Thể loại)", [
        ("id", "INTEGER", True, False),
        ("code", "TEXT", False, False),
        ("name", "TEXT", False, False),
        ("description", "TEXT", False, False)
    ], header_bg=COLOR_CYAN)

    # Dòng dưới (4 Bảng CSDL)
    draw_erd_entity(s7, 0.8, 4.0, 3.7, 2.4, "📚 BOOKS (Kho sách)", [
        ("id", "INTEGER", True, False),
        ("book_code", "TEXT", False, False),
        ("title / author", "TEXT", False, False),
        ("category_id", "INTEGER", False, True),
        ("total_qty / avail", "INT", False, False),
        ("rack_location", "TEXT", False, False)
    ], header_bg=COLOR_PRIMARY)

    draw_erd_entity(s7, 4.8, 4.0, 3.7, 2.4, "🤝 BORROW_RECORDS (Mượn/Trả)", [
        ("id", "INTEGER", True, False),
        ("borrow_code", "TEXT", False, False),
        ("reader_id", "INTEGER", False, True),
        ("book_id", "INTEGER", False, True),
        ("borrow_date / due", "DATE", False, False),
        ("fine_amount", "REAL", False, False)
    ], header_bg=COLOR_ROSE)

    draw_erd_entity(s7, 8.8, 4.0, 3.7, 1.15, "📌 RESERVATIONS (Đặt trước)", [
        ("id", "INTEGER", True, False),
        ("book_id / reader_id", "INT", False, True),
        ("queue_order", "INT", False, False)
    ], header_bg=COLOR_AMBER)

    draw_erd_entity(s7, 8.8, 5.25, 3.7, 1.15, "💰 FINE_LOGS (Thu phạt)", [
        ("id", "INTEGER", True, False),
        ("borrow_id / reader_id", "INT", False, True),
        ("amount / method", "TEXT", False, False)
    ], header_bg=COLOR_ROSE)

    # Thanh Banner chú thích Mối quan hệ giữa các Bảng phía dưới cùng
    rel_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.65))
    rel_card.fill.solid()
    rel_card.fill.fore_color.rgb = RGBColor(238, 242, 255) # Indigo nhẹ
    rel_card.line.color.rgb = COLOR_PRIMARY
    rel_card.line.width = Pt(1)

    tf_rel = rel_card.text_frame
    tf_rel.word_wrap = True
    tf_rel.margin_top = Inches(0.08)
    tf_rel.margin_left = Inches(0.2)
    p_rel0 = tf_rel.paragraphs[0]
    p_rel0.text = "🔗 CÁC MỐI QUAN HỆ CHÍNH GIỮA CÁC BẢNG DỮ LIỆU (DATABASE RELATIONSHIPS):"
    p_rel0.font.name = 'Segoe UI'
    p_rel0.font.size = Pt(10)
    p_rel0.font.bold = True
    p_rel0.font.color.rgb = COLOR_PRIMARY

    p_rel1 = tf_rel.add_paragraph()
    p_rel1.text = "• users.reader_id ──> readers.id (1-1)  |  • books.category_id ──> categories.id (N-1)  |  • borrow_records (reader_id, book_id) ──> (readers, books) (N-1)  |  • fine_logs.borrow_id ──> borrow_records.id (1-1)"
    p_rel1.font.name = 'Segoe UI'
    p_rel1.font.size = Pt(9)
    p_rel1.font.color.rgb = COLOR_TEXT_MAIN

    # ----------------------------------------------------
    # SLIDE 8: Kiến trúc Hệ thống & Giao tiếp
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_background(s8)
    add_header(s8, "Slide 8: Kiến Trúc Hệ Thống & Mô Hình Giao Tiếp")

    add_card(s8, 0.8, 1.5, 3.7, 5.5, "💻 Frontend Layer", [
        "• Single Page Application (SPA)",
        "• HTML5 Semantic Structure",
        "• Vanilla CSS3 Modern Light Theme",
        "• JavaScript ES6+ Fetch API",
        "• Thư viện Chart.js (Biểu đồ)"
    ], title_color=COLOR_CYAN, border_color=RGBColor(186, 230, 253), font_size=Pt(11.5))

    add_card(s8, 4.8, 1.5, 3.7, 5.5, "⚙️ Backend & DB Layer", [
        "• Python Custom HTTP Server (server.py)",
        "• RESTful API Endpoints (/api/...)",
        "• SQLite3 Database File (library.db)",
        "• SHA-256 Password Hashing",
        "• Auto Overdue Fine Calculator"
    ], title_color=COLOR_PRIMARY, border_color=RGBColor(199, 210, 254), font_size=Pt(11.5))

    add_card(s8, 8.8, 1.5, 3.7, 5.5, "🔄 Cách Giao Tiếp Bằng REST API", [
        "1. Người dùng thao tác trên Web SPA.",
        "2. Frontend gửi HTTP Request (GET/POST/PUT/DELETE) chứa dữ liệu JSON.",
        "3. Backend tiếp nhận, xử lý logic & truy vấn SQLite CSDL.",
        "4. Server trả kết quả JSON về cho Frontend hiển thị tức thì."
    ], title_color=COLOR_AMBER, border_color=RGBColor(254, 235, 200), font_size=Pt(11.5))

    # ----------------------------------------------------
    # SLIDE 9: Web Demo & Quy trình Demo đề xuất
    # ----------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_background(s9)
    add_header(s9, "Slide 9: Web Demo & Kịch Bản Demo Trình Diễn")

    add_card(s9, 0.8, 1.5, 5.7, 5.5, "🖥️ Các Màn Hình Mẫu Quan Trọng", [
        "1. Cổng Đăng Nhập (Login Gate): Sử dụng nút preset Đăng nhập nhanh.",
        "2. Trang Dashboard Thống Kê: Xem 4 thẻ KPIs & 2 Biểu đồ Chart.js.",
        "3. Trang Quản Lý Kho Sách: Xem danh sách dạng Grid/Table, thêm sách mới.",
        "4. Trang Quản Lý Độc Giả: Cấp mới thẻ độc giả & quản lý thời hạn.",
        "5. Trang Quản Lý Mượn/Trả/Phạt: Lập phiếu mượn, gia hạn & trả sách.",
        "6. Trang Đặt Trước: Xem hàng chờ đăng ký sách hết kho."
    ], title_color=COLOR_CYAN, border_color=RGBColor(186, 230, 253), font_size=Pt(11))

    add_card(s9, 6.8, 1.5, 5.7, 5.5, "📋 Thứ Tự Kịch Bản Demo Đề Xuất", [
        "Bước 1: Mở Web & Đăng nhập tài khoản Thủ thư (Preset: thuthu1 / 123456).",
        "Bước 2: Vào màn hình Tra cứu Sách -> Tìm kiếm từ khóa & Lọc thể loại.",
        "Bước 3: Thực hiện Thêm 1 cuốn sách mới vào kho -> Kiểm tra vị trí kệ.",
        "Bước 4: Sang tab Quản lý Độc giả -> Đăng ký 1 thẻ độc giả mới.",
        "Bước 5: Lập 1 Phiếu mượn sách mới -> Kiểm tra tồn kho giảm đi 1.",
        "Bước 6: Thực hiện Trả sách -> Kiểm tra tính tiền phạt nếu trễ hạn & Thu phạt.",
        "Bước 7: Xem Báo cáo Dashboard Thống kê tự động cập nhật."
    ], title_color=COLOR_EMERALD, border_color=RGBColor(167, 243, 208), font_size=Pt(11))

    # ----------------------------------------------------
    # SLIDE 10: Kết quả đạt được & Hướng phát triển
    # ----------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_background(s10)
    
    card10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3))
    card10.fill.solid()
    card10.fill.fore_color.rgb = COLOR_CARD
    card10.line.color.rgb = COLOR_EMERALD
    card10.line.width = Pt(2.5)

    tb10 = s10.shapes.add_textbox(Inches(1.2), Inches(0.9), Inches(10.933), Inches(5.7))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "KẾT QUẢ ĐẠT ĐƯỢC, HẠN CHẾ & HƯỚNG PHÁT TRIỂN"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLOR_EMERALD
    p.space_after = Pt(12)

    p = tf10.add_paragraph()
    p.text = "🎯 KẾT QUẢ ĐÃ HOÀN THÀNH:"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(3)

    res_bullets = [
        "Hoàn thiện 100% ứng dụng Web Quản lý Thư viện SPA mượt mà, màu sắc tươi sáng, tương thích mọi thiết bị.",
        "Thực hiện thành công toàn bộ nghiệp vụ: Đăng nhập/Phân quyền, Quản lý kho sách, Độc giả, Mượn/Trả/Phạt, Đặt trước & Dashboard thống kê."
    ]
    for b in res_bullets:
        p = tf10.add_paragraph()
        p.text = "  • " + b
        p.font.name = 'Segoe UI'
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(2)

    p = tf10.add_paragraph()
    p.text = "⚠️ HẠN CHẾ HIỆN TẠI:"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_ROSE
    p.space_after = Pt(3)

    lim_bullets = [
        "Chưa tích hợp phần cứng quét mã vạch Barcode/QR Code thực tế.",
        "Chưa có tính năng gửi Email/SMS tự động nhắc hạn trả cho độc giả."
    ]
    for b in lim_bullets:
        p = tf10.add_paragraph()
        p.text = "  • " + b
        p.font.name = 'Segoe UI'
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(2)

    p = tf10.add_paragraph()
    p.text = "🚀 HƯỚNG PHÁT TRIỂN TRONG TƯƠNG LAI:"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER
    p.space_after = Pt(3)

    future_bullets = [
        "Tích hợp Camera quét mã vạch Barcode/QR Code cho mượn sách 1s.",
        "Phát triển Mobile App (iOS/Android) gửi thông báo đẩy nhắc hạn mượn."
    ]
    for b in future_bullets:
        p = tf10.add_paragraph()
        p.text = "  • " + b
        p.font.name = 'Segoe UI'
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(2)

    p = tf10.add_paragraph()
    p.text = "\n✨ CHÂN THÀNH CẢM ƠN THẦY CÔ VÀ CÁC BẠN ĐÃ LẮNG NGHE! ✨"
    p.font.name = 'Segoe UI'
    p.font.size = Pt(14)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = COLOR_PRIMARY

    output_path = "Bao_Cao_Tien_Do_Quan_Ly_Thu_Vien.pptx"
    try:
        prs.save(output_path)
    except Exception:
        output_path = "Bao_Cao_Tien_Do_Quan_Ly_Thu_Vien_v2.pptx"
        prs.save(output_path)
    print(f"Successfully generated visual ERD PowerPoint presentation at: {os.path.abspath(output_path)}")

if __name__ == '__main__':
    create_10_slide_presentation()
