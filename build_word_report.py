import sys
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

def set_cell_background(cell, fill_color):
    """Đặt màu nền cho cell trong bảng Word"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def create_times_new_roman_word_report():
    doc = Document()

    # Cấu hình lề trang 2.54 cm (1 inch = 1440 dxa = 72 pt = 1.0 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    FONT_FAMILY = 'Times New Roman'
    COLOR_BLACK = RGBColor(0, 0, 0) # Chữ đen hoàn toàn theo yêu cầu

    # --- TIÊU ĐỀ TRANG BÁO CÁO PHÂN TÍCH HỆ THỐNG ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(18)
    
    r_sub = title_p.add_run("BÁO CÁO PHÂN TÍCH HỆ THỐNG THÔNG TIN\n")
    r_sub.font.name = FONT_FAMILY
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_BLACK

    r_main = title_p.add_run("ĐỀ TÀI: HỆ THỐNG QUẢN LÝ THƯ VIỆN\n(SMARTLIBRARY SYSTEM)")
    r_main.font.name = FONT_FAMILY
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = COLOR_BLACK

    # Thông tin nhóm & Đề tài
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.space_after = Pt(24)
    
    r_info = info_p.add_run(
        "Môn học: Phân Tích & Thiết Kế Hệ Thống Thông Tin / Lập Trình Web\n"
        "Tên Nhóm: Nhóm 3 - Lớp Quản Lý Thư Viện\n"
        "Thành viên thực hiện: Trương Thị Hạnh, Nguyễn Hồng Ngọc, Vy Thị Hồng Ngọc\n"
        "Ngày báo cáo: Năm học 2025 - 2026"
    )
    r_info.font.name = FONT_FAMILY
    r_info.font.size = Pt(12)
    r_info.font.italic = True
    r_info.font.color.rgb = COLOR_BLACK

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_BLACK
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_BLACK
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = COLOR_BLACK
        return h

    def add_body(text, bold=False, italic=False, space_after=6, indent=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        run = p.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.size = Pt(12)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = COLOR_BLACK
        return p

    def add_bullet(text, bold_prefix="", space_after=4):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = FONT_FAMILY
            r_bold.font.size = Pt(12)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_BLACK
        run = p.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_BLACK
        return p

    # ====================================================
    # PHẦN PHÂN TÍCH HỆ THỐNG CHI TIẾT
    # ====================================================

    # 1. TÊN ĐỀ TÀI & GIỚI THIỆU
    add_h1("1. TÊN ĐỀ TÀI VÀ MỤC TIÊU CỦA HỆ THỐNG")
    add_body("Tên đề tài: HỆ THỐNG QUẢN LÝ THƯ VIỆN (SMARTLIBRARY SYSTEM)", bold=True, indent=True)
    add_body(
        "Trong kỷ nguyên chuyển đổi số hiện nay, các tổ chức giáo dục và doanh nghiệp đang ngày càng đẩy mạnh việc ứng dụng công nghệ thông tin vào công tác quản lý tài nguyên tri thức. Thư viện truyền thống với phương thức quản lý bằng sổ sách ghi chép bộc lộ nhiều điểm hạn chế như tốn thời gian tra cứu, dễ nhầm lẫn hạn trả sách và khó kiểm soát số lượng kho thực tế. Hệ thống Quản lý Thư viện (SmartLibrary System) được xây dựng nhằm giải quyết triệt để các vấn đề trên, mang lại một giải pháp Web quản lý tập trung, hiện đại và minh bạch.",
        indent=True
    )

    add_h2("Mục tiêu cụ thể của hệ thống bao gồm:")
    add_bullet(" Chuyển đổi toàn bộ thông tin quản lý thủ công (sách, thể loại, thẻ độc giả, phiếu mượn/trả) sang cơ sở dữ liệu quan hệ tập trung SQLite3.", "Số hóa dữ liệu thư viện:")
    add_bullet(" Tự động hóa hoàn toàn việc trừ/cộng số lượng sách tồn kho khi mượn trả; tự động tính ngày hết hạn (14 ngày); và tự động quét tính tiền phạt quá hạn (5.000 VNĐ / ngày trễ).", "Tự động hóa nghiệp vụ:")
    add_bullet(" Cung cấp giao diện Web tra cứu 24/7 giúp độc giả biết chính xác vị trí kệ kho (vd: Kệ A1, B2) và chủ động đặt trước khi sách trong kho hết.", "Nâng cao trải nghiệm độc giả:")
    add_bullet(" Cung cấp công cụ báo cáo thống kê trực quan bằng biểu đồ cho Quản trị viên và Thủ thư theo dõi các chỉ số hoạt động realtime.", "Hỗ trợ quản trị & thống kê:")

    # 2. ĐỐI TƯỢNG SỬ DỤNG
    add_h1("2. ĐỐI TƯỢNG SỬ DỤNG HỆ THỐNG (3 VAI TRÒ - RBAC)")
    add_body("Hệ thống được thiết kế phân quyền truy cập chặt chẽ thành 3 nhóm đối tượng người dùng chính:", indent=True)

    add_bullet("Là người có quyền hạn cao nhất trong hệ thống. Admin có nhiệm vụ khởi tạo và quản lý toàn bộ tài khoản người dùng (users), cấp vai trò (admin, librarian, reader), đặt lại mật khẩu và xem Bảng điều khiển thống kê tổng quan (Dashboard KPIs).", "1. Quản trị viên (Admin System):")
    add_bullet("Là cán bộ vận hành trực tiếp tại thư viện. Thủ thư có quyền quản lý danh mục kho sách (thêm mới, chỉnh sửa thông tin, đổi vị trí kệ, điều chỉnh số lượng tồn); quản lý hồ sơ độc giả (cấp mới thẻ Sinh viên/Giảng viên, khóa/mở thẻ); lập phiếu mượn/trả sách tại quầy; xử lý gia hạn mượn và thực hiện thu tiền phạt quá hạn.", "2. Thủ thư (Librarian Operations):")
    add_bullet("Là đối tượng học sinh, sinh viên, giảng viên hoặc độc giả ngoài sử dụng thư viện. Độc giả sử dụng hệ thống để tra cứu thông tin sách trực tuyến, xem vị trí kệ kho chính xác, theo dõi các phiếu mượn cá nhân và đăng ký xếp hàng Đặt trước khi sách bị mượn hết.", "3. Độc giả (Reader Experience):")

    # 3. CÁC VẤN ĐỀ MÀ HỆ THỐNG GIẢI QUYẾT
    add_h1("3. CÁC VẤN ĐỀ MÀ HỆ THỐNG GIẢI QUYẾT")
    add_body("Hệ thống Quản lý Thư viện giải quyết trực tiếp 4 bài toán thực tế nghiêm trọng của các thư viện truyền thống:", indent=True)

    add_bullet("Thư viện truyền thống bắt độc giả tìm kiếm qua tủ thẻ hoặc sổ sách thủ công. Hệ thống mới cung cấp bộ lọc tìm kiếm thời gian thực (Real-time Filter) theo Tên sách, Tác giả, NXB và hiển thị chính xác vị trí Kệ kho (ví dụ: Kệ A1-01, Kệ B2-05) giúp độc giả đến tận nơi lấy sách trong vài giây.", "1. Giải quyết bài toán tra cứu chậm và không biết vị trí sách:")
    add_bullet("Ghi chép sổ sách thủ công dễ dẫn đến sai sót khi tính ngày quá hạn. Hệ thống tự động so sánh Ngày trả thực tế với Hạn trả mặc định (14 ngày), nếu trễ hạn sẽ tự động nhân 5.000 VNĐ / ngày trễ, đồng thời ghi nhận hóa đơn thu phạt minh bạch.", "2. Khắc phục sai sót khi tính tiền phạt trễ hạn:")
    add_bullet("Hệ thống duy trì đồng bộ 2 chỉ số: Tổng số lượng sách nhập kho (total_qty) và Số lượng còn sẵn có thực tế (available_qty). Khi lập phiếu mượn, available_qty tự động trừ 1; khi trả sách, available_qty tự động cộng 1, giúp thủ thư kiểm soát chính xác lượng sách tại thời điểm mượn.", "3. Loại bỏ rủi ro thất thoát kho sách:")
    add_bullet("Khi một cuốn sách hay bị mượn hết (available_qty = 0), hệ thống cung cấp nút 'Đặt trước' giúp độc giả đăng ký xếp hàng ưu tiên mượn (queue_order) ngay khi cuốn sách được người trước mang trả.", "4. Giải quyết tình trạng mượn chồng chéo khi hết sách:")

    # 4. YÊU CẦU CHỨC NĂNG & PHI CHỨC NĂNG
    add_h1("4. PHÂN TÍCH YÊU CẦU HỆ THỐNG")

    add_h2("4.1. Yêu cầu chức năng (Functional Requirements)")
    add_bullet("Giao diện phủ màn hình Login Gate Overlay bắt buộc. Xác thực tên đăng nhập/mật khẩu với mã hóa SHA-256. Phân quyền truy cập 3 nhóm vai trò (Admin, Thủ thư, Độc giả). Hỗ trợ nút Preset Accounts đăng nhập nhanh dùng thử.", "Yêu cầu Đăng nhập & Bảo mật:")
    add_bullet("Cho phép xem danh mục sách ở dạng Grid thẻ hoặc Bảng Table. Tìm kiếm từ khóa tức thì theo tên sách, tác giả, mã sách. Lọc sách theo 6 danh mục thể loại. Thêm mới, chỉnh sửa thông tin sách, đổi vị trí kệ kho và xóa sách.", "Yêu cầu Quản lý Sách & Thể loại:")
    add_bullet("Quản lý thông tin độc giả (Mã độc giả, Họ tên, Email, SĐT, Loại thẻ, Trạng thái). Cấp mới thẻ độc giả với quy định hạn thẻ tự động (Sinh viên: 2 năm, Giảng viên: 4 năm). Khóa/Mở thẻ độc giả.", "Yêu cầu Quản lý Độc giả:")
    add_bullet("Lập phiếu mượn mới (kiểm tra thẻ còn hạn & tồn kho > 0). Hạn trả mặc định 14 ngày. Gia hạn mượn +14 ngày (tối đa 2 lần). Tiếp nhận trả sách, tự tính tiền phạt 5.000đ/ngày trễ & Lập hóa đơn thu tiền phạt.", "Yêu cầu Mượn / Trả / Phạt:")
    add_bullet("Kích hoạt đăng ký khi available_qty = 0, tự động cấp số thứ tự xếp hàng queue_order.", "Yêu cầu Đặt trước Sách:")
    add_bullet("Hiển thị 4 thẻ KPIs Realtime (Tổng sách, Độc giả, Phiếu mượn/Quá hạn, Tiền phạt) và 2 Biểu đồ Chart.js (Cơ cấu thể loại & Lưu lượng mượn/trả).", "Yêu cầu Thống kê & Báo cáo:")

    add_h2("4.2. Yêu cầu phi chức năng (Non-Functional Requirements)")
    add_bullet("Giao diện Trang đơn (SPA) phản hồi thao tác dưới 1 giây, tra cứu lọc từ khóa không cần tải lại trang.", "Hiệu năng & Tốc độ phản hồi:")
    add_bullet("Mật khẩu tài khoản được mã hóa bằng thuật toán SHA-256. Kiểm tra session đăng nhập và phân quyền API chặt chẽ.", "Bảo mật dữ liệu:")
    add_bullet("CSDL SQLite3 bật cấu hình PRAGMA foreign_keys = ON; đảm bảo toàn vẹn ràng buộc khóa ngoại giữa các bảng.", "Toàn vẹn dữ liệu:")
    add_bullet("Giao diện thiết kế theo chuẩn Modern Light Design tươi sáng, trực quan, dễ thao tác.", "Tính dễ sử dụng:")

    # 5. DANH SÁCH & MÔ TẢ CHI TIẾT CÁC CHỨC NĂNG HỆ THỐNG
    add_h1("5. MÔ TẢ CHI TIẾT 11 CHỨC NĂNG HIỆN CÓ IN PROJECT")
    add_body("Bảng dưới đây liệt kê chi tiết 11 chức năng thực tế đã hoàn thành 100% trong mã nguồn project:", indent=True)

    table_fn = doc.add_table(rows=1, cols=3)
    table_fn.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_fn.rows[0].cells
    hdr_cells[0].text = "STT"
    hdr_cells[1].text = "Tên Chức Năng"
    hdr_cells[2].text = "Mô Tả Chi Tiết Hoạt Động Trong Mã Nguồn"

    for cell in hdr_cells:
        set_cell_background(cell, "EAEAEA")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = FONT_FAMILY
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_BLACK

    fn_data = [
        ("1", "Cổng Đăng Nhập & Phân Quyền", "Màn hình phủ loginOverlay bắt buộc. Nhập username/password -> Xác thực API /api/auth/login -> Phân quyền 3 vai trò (admin, librarian, reader). Hỗ trợ 3 nút Preset Accounts đăng nhập dùng thử nhanh."),
        ("2", "Tra Cứu & Tìm Kiếm Sách", "Tìm kiếm real-time lọc theo Tên sách, Tác giả, Mã sách, NXB. Lọc theo danh mục thể loại. Chuyển đổi 2 chế độ hiển thị (Grid View / Table View). Hiển thị rõ cột Vị trí kệ kho (Kệ A1-01) và số lượng khả dụng."),
        ("3", "Quản Lý Kho Sách (Staff)", "Form Modal thêm/sửa sách: Nhập Mã sách, Tiêu đề, Tác giả, Thể loại, NXB, Năm XB, Số lượng tổng (total_qty), Số lượng sẵn có (available_qty), Vị trí kệ kho & URL ảnh bìa. Xóa sách khỏi kho."),
        ("4", "Quản Lý Thẻ Độc Giả (Staff)", "Bảng danh sách độc giả. Form cấp mới thẻ: Chọn loại thẻ (Sinh viên: hạn 2 năm, Giảng viên: hạn 4 năm). Tự động tính expiry_date. Chức năng Khóa / Kích hoạt thẻ độc giả."),
        ("5", "Lập Phiếu Mượn Sách Tại Quầy", "Form mượn: Chọn Độc giả + Chọn Sách -> Hệ thống kiểm tra thẻ còn hạn & tồn kho available_qty > 0 -> Tạo phiếu mượn với status = 'Đang mượn', Hạn trả mặc định 14 ngày, tự động trừ available_qty đi 1."),
        ("6", "Gia Hạn Hạn Mượn Sách", "Nút 'Gia hạn' trên phiếu mượn -> Kiểm tra số lần gia hạn (renewal_count < 2) -> Cộng thêm 14 ngày vào hạn trả due_date và cập nhật renewal_count = renewal_count + 1."),
        ("7", "Xác Nhận Trả Sách & Phạt", "Nút 'Trả sách' -> Ghi nhận return_date, cập nhật status = 'Đã trả', cộng trả lại available_qty thêm 1. So sánh ngày trả với hạn trả: Nếu trễ, tự động nhân 5.000đ * số ngày trễ -> Cập nhật fine_status = 'Chưa nộp'."),
        ("8", "Thu Tiền Phạt Quá Hạn", "Nút 'Thu tiền phạt' -> Chọn phương thức (Tiền mặt / Chuyển khoản) -> Cập nhật fine_status = 'Đã nộp' -> Thêm bản ghi mới vào bảng fine_logs."),
        ("9", "Đặt Trước Sách (Reservations)", "Khi độc giả muốn mượn sách có available_qty = 0, nhấn nút 'Đặt trước' -> Tạo bản ghi trong bảng reservations với queue_order tăng tự động để xếp hàng chờ mượn."),
        ("10", "Dashboard Thống Kê (Chart.js)", "4 Thẻ chỉ số KPIs Realtime (Tổng đầu sách/tổng cuốn, Độc giả đang hoạt động, Phiếu mượn đang lưu hành/quá hạn, Doanh thu phạt đã thu). Biểu đồ Doughnut Chart (thể loại) & Bar Chart (lượt mượn/trả)."),
        ("11", "Quản Lý Tài Khoản (Admin)", "Tab quản lý users dành riêng cho Admin: Thêm tài khoản mới, chọn vai trò (admin, librarian, reader), liên kết Mã độc giả, đổi mật khẩu.")
    ]

    for row_idx, data in enumerate(fn_data):
        row_cells = table_fn.add_row().cells
        row_cells[0].text = data[0]
        row_cells[1].text = data[1]
        row_cells[2].text = data[2]
        bg_color = "F9F9F9" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, cell in enumerate(row_cells):
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = FONT_FAMILY
                run.font.size = Pt(10.5)
                run.font.color.rgb = COLOR_BLACK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 6. QUY TRÌNH HOẠT ĐỘNG CỦA HỆ THỐNG
    add_h1("6. CÁC QUY TRÌNH HOẠT ĐỘNG NGHIỆP VỤ HỆ THỐNG")

    add_h2("6.1. Quy trình Quản lý Sách")
    add_bullet("Thủ thư đăng nhập tài khoản và chọn Tab 'Tra cứu & Quản lý Sách'.", "Bước 1: ")
    add_bullet("Nhấn nút 'Thêm Sách Mới' để mở cửa sổ Modal nhập liệu.", "Bước 2: ")
    add_bullet("Điền đầy đủ thông tin: Mã sách, Tiêu đề, Tác giả, Chọn Thể loại, NXB, Năm XB, Số lượng tổng (total_qty), Vị trí kệ kho (như Kệ A1-01) và Ảnh bìa URL.", "Bước 3: ")
    add_bullet("Nhấn 'Lưu Sách'. Hệ thống gửi yêu cầu API POST /api/books về máy chủ Python server.py.", "Bước 4: ")
    add_bullet("Máy chủ lưu thông tin vào CSDL SQLite library.db, tự động gán available_qty = total_qty và trả về kết quả thành công. Giao diện Web SPA tự động hiển thị sách mới mà không cần tải lại trang.", "Bước 5: ")

    add_h2("6.2. Quy trình Quản lý Người dùng / Độc giả")
    add_bullet("Thủ thư mở Tab 'Quản lý Độc giả' và nhấn 'Cấp Thẻ Độc Giả Mới'.", "Bước 1: ")
    add_bullet("Nhập thông tin cá nhân độc giả: Họ tên, Email, Số điện thoại và Chọn Loại thẻ (Sinh viên, Giảng viên, Độc giả ngoài).", "Bước 2: ")
    add_bullet("Nhấn 'Lưu Độc giả'. Máy chủ tự động cấp Mã độc giả (DG001...) và tính Ngày hết hạn thẻ (Sinh viên: +2 năm, Giảng viên: +4 năm).", "Bước 3: ")
    add_bullet("Khi thẻ bị quá hạn hoặc độc giả vi phạm quy định mượn sách, Thủ thư có thể nhấn nút 'Khóa thẻ'. Hệ thống chuyển status = 'Bị khóa', ngăn không cho lập phiếu mượn mới.", "Bước 4: ")

    add_h2("6.3. Quy trình Cho Mượn Sách Tại Quầy")
    add_bullet("Thủ thư mở Tab 'Quản lý Mượn/Trả/Phạt' và nhấn 'Tạo Phiếu Mượn Sách'.", "Bước 1: ")
    add_bullet("Chọn Độc giả mượn và Chọn cuốn sách cần mượn từ danh mục.", "Bước 2: ")
    add_bullet("Hệ thống tự động kiểm tra các điều kiện: Thẻ độc giả còn hạn hoạt động? Số lượng khả dụng của sách có lớn hơn 0 (available_qty > 0)?", "Bước 3: ")
    add_bullet("Nếu hợp lệ, hệ thống tạo bản ghi phiếu mượn mới trong bảng borrow_records với Hạn trả mặc định 14 ngày, đồng thời tự động giảm available_qty của cuốn sách đi 1.", "Bước 4: ")
    add_bullet("Trong quá trình mượn, nếu độc giả chưa đọc xong, Thủ thư hoặc Độc giả có thể bấm 'Gia hạn'. Hệ thống cộng thêm 14 ngày vào hạn trả (tối đa 2 lần gia hạn).", "Bước 5: ")

    add_h2("6.4. Quy trình Trả Sách & Xử Lý Tiền Phạt Quá Hạn")
    add_bullet("Độc giả mang sách đến trả tại quầy. Thủ thư tìm phiếu mượn và nhấn nút 'Xác nhận Trả sách'.", "Bước 1: ")
    add_bullet("Hệ thống ghi nhận Ngày trả thực tế (return_date), cập nhật status = 'Đã trả', và cộng trả lại available_qty thêm 1 cho kho sách.", "Bước 2: ")
    add_bullet("Máy tính tự động so sánh Ngày trả thực tế với Hạn trả mặc định (due_date). Nếu trễ hạn, hệ thống nhân 5.000 VNĐ / mỗi ngày trễ, cập nhật fine_amount và gán fine_status = 'Chưa nộp'.", "Bước 3: ")
    add_bullet("Thủ thư nhấn nút 'Thu tiền phạt', chọn phương thức (Tiền mặt hoặc Chuyển khoản). Hệ thống cập nhật fine_status = 'Đã nộp' và lưu một bản ghi nhật ký mới vào bảng fine_logs để quản lý doanh thu.", "Bước 4: ")

    # 7. PHÂN TÍCH CƠ SỞ DỮ LIỆU CHI TIẾT (7 BẢNG)
    add_h1("7. PHÂN TÍCH CƠ SỞ DỮ LIỆU (7 BẢNG CSDL SQLITE3)")
    add_body("Cơ sở dữ liệu của dự án được lưu trữ trong file library.db bao gồm 7 bảng dữ liệu quan hệ được chuẩn hóa nghiêm ngặt:", indent=True)

    add_bullet("id (INTEGER, PK, Auto), username (TEXT, UNIQUE), password_hash (TEXT, SHA-256), role (TEXT: admin/librarian/reader), full_name (TEXT), email (TEXT), reader_id (INTEGER, FK). Bảng dùng để xác thực và phân quyền người dùng.", "1. Bảng users (Tài khoản người dùng): ")
    add_bullet("id (INTEGER, PK, Auto), reader_code (TEXT, UNIQUE), full_name (TEXT), email (TEXT), phone (TEXT), card_type (TEXT: Sinh viên/Giảng viên/Độc giả ngoài), status (TEXT: Hoạt động/Bị khóa), issue_date (DATE), expiry_date (DATE).", "2. Bảng readers (Thông tin thẻ độc giả): ")
    add_bullet("id (INTEGER, PK, Auto), code (TEXT, UNIQUE), name (TEXT), description (TEXT). Lưu 6 danh mục thể loại (CNTT, Văn học, Kinh tế, Khoa học, Lịch sử, Kỹ năng sống).", "3. Bảng categories (Thể loại sách): ")
    add_bullet("id (INTEGER, PK, Auto), book_code (TEXT, UNIQUE), title (TEXT), author (TEXT), category_id (INTEGER, FK), publisher (TEXT), publish_year (INTEGER), total_qty (INTEGER), available_qty (INTEGER), rack_location (TEXT), description (TEXT), cover_url (TEXT). Bảng quản lý kho sách và vị trí kệ.", "4. Bảng books (Kho sách thư viện): ")
    add_bullet("id (INTEGER, PK, Auto), borrow_code (TEXT, UNIQUE), reader_id (INTEGER, FK), book_id (INTEGER, FK), borrow_date (DATE), due_date (DATE), return_date (DATE), status (TEXT: Đang mượn/Đã trả/Quá hạn), renewal_count (INTEGER), fine_amount (REAL), fine_status (TEXT: Chưa nộp/Đã nộp).", "5. Bảng borrow_records (Phiếu mượn/trả/phạt): ")
    add_bullet("id (INTEGER, PK, Auto), book_id (INTEGER, FK), reader_id (INTEGER, FK), request_date (DATETIME), status (TEXT: Đang chờ/Đã duyệt), queue_order (INTEGER). Hàng chờ đặt trước khi sách hết kho.", "6. Bảng reservations (Đặt trước sách): ")
    add_bullet("id (INTEGER, PK, Auto), borrow_id (INTEGER, FK), reader_id (INTEGER, FK), amount (REAL), paid_at (DATETIME), payment_method (TEXT: Tiền mặt/Chuyển khoản). Lịch sử thu tiền phạt quá hạn.", "7. Bảng fine_logs (Nhật ký đóng tiền phạt): ")

    add_h2("Mối quan hệ chính giữa các bảng (Database Relationships):")
    add_bullet("Một tài khoản người dùng có thể liên kết với một hồ sơ độc giả tương ứng.", "users.reader_id ──> readers.id (Quan hệ 1 - 1): ")
    add_bullet("Một danh mục thể loại có thể chứa nhiều đầu sách khác nhau.", "books.category_id ──> categories.id (Quan hệ N - 1): ")
    add_bullet("Một độc giả có thể lập nhiều phiếu mượn sách theo thời gian.", "borrow_records.reader_id ──> readers.id (Quan hệ N - 1): ")
    add_bullet("Một đầu sách có thể xuất hiện trong nhiều phiếu mượn của các độc giả khác nhau.", "borrow_records.book_id ──> books.id (Quan hệ N - 1): ")
    add_bullet("Mỗi phiếu mượn bị phạt quá hạn sẽ phát sinh một bản ghi lịch sử nộp phạt.", "fine_logs.borrow_id ──> borrow_records.id (Quan hệ 1 - 1): ")

    # 8. KIẾN TRÚC HỆ THỐNG & CÔNG NGHỆ SỬ DỤNG
    add_h1("8. KIẾN TRÚC HỆ THỐNG VÀ TẬP CÔNG NGỆ SỬ DỤNG")
    add_body("Hệ thống được thiết kế theo mô hình Kiến trúc RESTful Client-Server decouple hiện đại:", indent=True)

    add_bullet("Ứng dụng Web Trang đơn (Single Page Application - SPA) được xây dựng hoàn toàn bằng HTML5 ngữ nghĩa, Vanilla CSS3 (Modern Light Design) và JavaScript ES6+. Sử dụng Fetch API để gửi nhận dữ liệu không đồng bộ và thư viện Chart.js để vẽ biểu đồ thống kê trực quan.", "1. Tầng Giao diện Người dùng (Frontend):")
    add_bullet("Máy chủ HTTP được phát triển tùy biến bằng ngôn ngữ Python (server.py) lắng nghe tại cổng 8000. Máy chủ chịu trách nhiệm tiếp nhận các yêu cầu HTTP (GET, POST, PUT, DELETE), thực hiện logic nghiệp vụ kiểm tra thẻ, kiểm tra tồn kho và định kỳ quét tính trễ hạn.", "2. Tầng Máy chủ Xử lý (Backend):")
    add_bullet("Cơ sở dữ liệu SQLite3 (library.db) lưu trữ dữ liệu tập trung trong 1 file local. Bật cấu hình PRAGMA foreign_keys = ON; để đảm bảo ràng buộc toàn vẹn dữ liệu. Mật khẩu tài khoản được mã hóa bằng thuật toán SHA-256.", "3. Tầng Cơ sở Dữ liệu (Database):")

    # 9. ĐÁNH GIÁ ƯU ĐIỂM, HẠN CHẾ & HƯỚNG PHÁT TRIỂN
    add_h1("9. ĐÁNH GIÁ ƯU ĐIỂM, HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN")

    add_h2("9.1. Ưu điểm nổi bật")
    add_bullet("Giao diện Light Mode tươi sáng, hiện đại, hỗ trợ phản hồi mượt mà dưới 1 giây không cần tải lại trang (SPA).")
    add_bullet("Hiển thị minh bạch cột vị trí kệ kho (Kệ A1, B2) giúp độc giả và thủ thư tìm thấy sách ngay lập tức.")
    add_bullet("Tự động hóa 100% các nghiệp vụ tính toán phức tạp: Trừ/cộng tồn kho, đặt hạn mượn, gia hạn và tính phạt 5.000đ/ngày trễ.")
    add_bullet("Phân quyền 3 vai trò chặt chẽ, bảo mật mật khẩu SHA-256.")

    add_h2("9.2. Hạn chế hiện tại")
    add_bullet("Chưa kết nối với thiết bị phần cứng quét mã vạch Barcode / QR Code thực tế tại quầy.")
    add_bullet("Chưa có dịch vụ gửi Email hoặc tin nhắn SMS tự động nhắc nợ khi sách sắp đến hạn trả.")

    add_h2("9.3. Hướng phát triển trong tương lai")
    add_bullet("Tích hợp Camera máy tính/điện thoại quét mã vạch Barcode trên bìa sách để lập phiếu mượn trong 1 giây.")
    add_bullet("Phát triển ứng dụng di động Mobile App (iOS / Android) cho độc giả tra cứu và nhận thông báo nhắc hạn mượn.")
    add_bullet("Mở rộng tính năng mượn liên thư viện giữa nhiều chi nhánh.")

    output_path = "Phan_Tich_He_Thong_Quan_Ly_Thu_Vien.docx"
    try:
        doc.save(output_path)
    except Exception:
        output_path = "Phan_Tich_He_Thong_Quan_Ly_Thu_Vien_v2.docx"
        doc.save(output_path)
    print(f"Successfully generated Times New Roman Word report at: {os.path.abspath(output_path)}")

if __name__ == '__main__':
    create_times_new_roman_word_report()
